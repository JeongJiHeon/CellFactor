#!/usr/bin/env python3
"""보충 표 묶음(Supplementary Tables S1–S10) — 저장 산출만 읽음(측정 비율만 원자료에서 직접 계산).
출력 final_results/figures/manuscript_v2/Supp/Supplementary_Tables.xlsx (+ 시트별 csv: Supp/tables/)"""
import os
import json, sys, gzip
from pathlib import Path
import numpy as np, pandas as pd, torch
ROOT = Path(__file__).resolve().parents[4]; FR = ROOT / "final_results"; MR = ROOT / "results/mech"; D0 = Path(os.environ.get("CELLFACTOR_DATA", "data"))
OUT = FR / "figures/manuscript_v2/Supp"; (OUT / "tables").mkdir(parents=True, exist_ok=True)
ORG = [("breast", "Breast"), ("liver", "Liver"), ("pancreas", "Pancreas")]
S = {}

# ---------- S1 C-index per cohort–endpoint ----------
rows = []
for o, on in ORG:
    R = json.load(open(FR / f"02_cindex/bootstrap_ci_B2000/{o}_BOOT.json"))
    for m, d in R["models"].items():
        if m == "CellFactor (c1–c6)": continue   # 2026-10-11: CellFactor 자체가 c1–c6 이라 중복 행
        if m == "State-only model": continue   # CellFactor 절제판 → S9
        for k, v in d.items():
            scope = "held-out mean" if k == "heldout" else ("external mean" if k == "external_mean" else ("held-out" if k.startswith("TCGA") else ("external (primary)" if k in R["external_endpoints"] else "external (appendix)")))
            if scope == "held-out": continue
            pd6 = R.get("paired_vs_CellFactor_c1_6", {}).get(m, {}).get(k)   # 2026-10-10: 짝 차이 = CellFactor (c1–c6) − 모델, 같은 재표집
            rows.append({"Cancer": on, "Set": scope, "Cohort|endpoint": k if "|" in k else "", "Model": m, "C-index": v["point"], "CI low": v["ci95"][0], "CI high": v["ci95"][1], "Bootstrap SE": v["se"],
                         "Paired diff, CellFactor (c1–c6) − model": pd6["diff"] if pd6 else None, "Diff CI low": pd6["ci95"][0] if pd6 else None,
                         "Diff CI high": pd6["ci95"][1] if pd6 else None, "Diff P (two-sided bootstrap)": pd6["p_two_sided"] if pd6 else None})
S["S1_Cindex_by_endpoint"] = pd.DataFrame(rows)
cpc = pd.read_csv(FR / "02_cindex/cindex_per_cohort.csv")

# ---------- S2 cohorts and datasets ----------
def panel(o): return [g.strip() for g in open(ROOT / f"results/tcga_panel/{o}/cand3k.txt") if g.strip()]
cov = {}
b = torch.load(D0 / "external_bulk/scanb_16k/SCANB_BUNDLE.pt", weights_only=False); P = set(panel("breast"))
cov["SCAN-B"] = np.mean([bool(p) for g, p in zip(b["gene_names"], b["gene_present"]) if g in P])
mp = np.load(D0 / "external_bulk/metabric_16k/METABRIC_PRESENT_16K.npy"); cov["METABRIC"] = np.mean([bool(p) for g, p in zip(b["gene_names"], mp) if g in P])
cov["GSE14520"] = cov["GSE76427"] = 1.0
sys.path.insert(0, str(ROOT / "code/lib")); import panc_stage23 as PS
pp = torch.load(ROOT / "sandbox6/results/pancreas/PANC_STAGE1_DATA.pt", map_location="cpu", weights_only=False, mmap=True)
genes = [str(g) for g in pp["gene_names"]]; Pp = set(panel("pancreas"))
for acc in PS.EXT:
    _, ob, _, _ = PS.load_ext(acc, genes); cov[acc] = float(np.mean([o_ for g, o_ in zip(genes, ob) if g in Pp]))
PLAT = {"SCAN-B": "RNA-seq (Illumina)", "METABRIC": "Microarray (Illumina HT-12)", "GSE14520": "Microarray (Affymetrix HT-U133A, GPL3921)",
        "GSE76427": "Microarray (Illumina HT-12, GPL10558)", "GSE62452": "Microarray (Affymetrix HuGene 1.0 ST, GPL6244)",
        "GSE57495": "Microarray (Rosetta/Merck custom, GPL15048)", "GSE85916": "Microarray (Affymetrix HG-U219, GPL13667)", "GSE10143": "Microarray (DASL, GPL5474)", "GSE54236": "Microarray (Agilent, GPL6480)"}
rows = []
for o, on in ORG:
    C = pd.read_csv(FR / f"01_umap/{o}/cells.csv.gz", usecols=["cell_type", "donor"])
    rows.append({"Cancer": on, "Role": "Single-cell reference (displayed cells)", "Dataset": "", "Platform": "scRNA-seq", "Patients/donors": C.donor.nunique(), "Cells": len(C), "Cell types": C.cell_type.nunique(), "Endpoint": "", "Events": "", "Panel coverage": ""})
    pb = torch.load(ROOT / f"results/hvg_design/run1a/{o}/A_cand3k/c1/pred_bulk.pt", map_location="cpu", weights_only=False)
    ev = pb["event_all"].numpy(); tr, ho = pb["train_idx"].numpy(), pb["heldout_idx"].numpy()
    rows.append({"Cancer": on, "Role": "TCGA training", "Dataset": f"TCGA-{ {'breast': 'BRCA', 'liver': 'LIHC', 'pancreas': 'PAAD'}[o]}", "Platform": "RNA-seq", "Patients/donors": len(tr), "Endpoint": "OS", "Events": int(ev[tr].sum()), "Panel coverage": 1.0})
    rows.append({"Cancer": on, "Role": "TCGA held-out", "Dataset": f"TCGA-{ {'breast': 'BRCA', 'liver': 'LIHC', 'pancreas': 'PAAD'}[o]}", "Platform": "RNA-seq", "Patients/donors": len(ho), "Endpoint": "OS", "Events": int(ev[ho].sum()), "Panel coverage": 1.0})
    for _, r in cpc[(cpc.organ == o) & (cpc.set != "held-out")].iterrows():
        rows.append({"Cancer": on, "Role": "External bulk (" + ("primary" if r.set == "external_primary" else "appendix") + ")", "Dataset": {"GSE10143": "GSE10143 (HCC tumor samples only)", "GSE54236": "GSE54236 (HCC tumor samples only)"}.get(r.cohort, r.cohort), "Platform": PLAT.get(r.cohort, ""),
                     "Patients/donors": int(r.n), "Endpoint": r.endpoint, "Events": int(r.events), "Panel coverage": round(cov[r.cohort], 3) if r.cohort in cov else ""})
SP = [("Breast", "TNBC (Wang et al. 2024; Zenodo 8135721; pathologist annotation)", "Legacy Spatial Transcriptomics array (1,934 spots, 100 µm, 150 µm center-to-center)"),
      ("Liver", "Wu et al. HCC (H&E annotation)", "10x CytAssist Visium"), ("Pancreas", "Cui Zhou et al. PDAC (cell-type annotation; HEST-1k)", "10x Visium")]
for on, ds, pf in SP: rows.append({"Cancer": on, "Role": "Spatial transcriptomics", "Dataset": ds, "Platform": pf})
S["S2_Cohorts_datasets"] = pd.DataFrame(rows)

# ---------- S3 program dictionary ----------
K0 = {"breast": ROOT / "data/registry/PANEL16K_PROGRAM_SETS_V2.json", "liver": ROOT / "results/liver/LIVER_PROGRAM_SETS.json", "pancreas": ROOT / "data/organ_aux/pancreas_6type/PANC_PROGRAM_SETS.json"}
rows = []
for o, on in ORG:
    k0 = json.load(open(K0[o])); note = {r["program"]: r for r in json.load(open(ROOT / f"results/tcga_panel/design/dicts/{o}_MIN_DICT_NOTE.json"))["lists"]["cand3k"]["rows"]}
    for p, gs in k0.items():
        gs = gs if isinstance(gs, list) else gs.get("genes", [])
        src = "MSigDB Hallmark" if p.startswith("HM_") else ("Curated common cell state (CS)" if p.startswith("CS_") else ("Curated cell-type/cancer-specific (markers from Öhlund et al. 2017; Elyada et al. 2019)" if p in ("CAF_MYO", "CAF_INFL") else "Curated cell-type/cancer-specific"))
        n = note.get(p, {})
        rows.append({"Cancer": on, "Program": p, "Source": src, "Genes (organ universe)": len(gs), "Genes in selected panel": n.get("n_list", ""),
                     "Panel-vs-full score r (within type)": n.get("r_within", ""), "Named axis in CellFactor": "yes" if n.get("kept") else "no",
                     "Reason if removed": n.get("reason") or "", "Genes": ", ".join(gs)})
S["S3_Program_dictionary"] = pd.DataFrame(rows)

# ---------- S4 selected gene panel ----------
cand = json.load(open(ROOT / "results_json/tcga_panel/design/CP_PCH_PANEL_PROGRAM_CANDIDATES_20260927.json"))["organs"]
rows = []
for o, on in ORG:
    kd = json.load(open(ROOT / f"results/hvg_design/plan_hvg/{o}_PROGRAM_SETS_K0DEDUP.json")); k0 = json.load(open(K0[o]))
    compact = {g for p, gs in kd.items() if not p.startswith("HM_") for g in (gs if isinstance(gs, list) else gs.get("genes", []))}
    seeds = {g for v in cand[o]["program_coverage"].values() for g in v.get("in_current_full", [])}
    marker = {}
    for t, v in cand[o]["marker_protection_by_type"].items():
        for g in v.get("selected_genes", []): marker.setdefault(g, []).append(t)
    below = set(cand[o]["protected_genes_below_generic_tcga_detection"])
    memb = {}
    for p, gs in k0.items():
        for g in (gs if isinstance(gs, list) else gs.get("genes", [])): memb.setdefault(g, []).append(p)
    for g in panel(o):
        cat = [c for c, ok in (("curated program gene", g in compact), ("curated identity/program core", g in seeds), ("reference cell-type marker", g in marker)) if ok]
        rows.append({"Cancer": on, "Gene": g, "Inclusion": "; ".join(cat) if cat else "TCGA-training variability rank",
                     "Marker cell types": ", ".join(marker.get(g, [])), "Program membership": ", ".join(memb.get(g, [])),
                     "Below generic TCGA detection (protected)": "yes" if g in below else ""})
S["S4_Selected_gene_panel"] = pd.DataFrame(rows)

# ---------- S5–S9 ----------
cox = json.load(open(MR / "exp3/COX_ADJ.json")); rows = []
for r in cox:
    base = {"Cancer": r["organ"].capitalize(), "Cohort": r["cohort"], "n": r["n"], "Events": r["events"], "Covariates": r["covariates"]}
    rows.append({**base, "Term": "CellFactor risk (per SD)", "HR unadjusted": r["HR_unadj"], "Unadj. CI low": r["HR_unadj_lo"], "Unadj. CI high": r["HR_unadj_hi"], "P unadjusted": r["p_unadj_wald"], "HR adjusted": r["HR_adj"], "CI low": r["HR_lo"], "CI high": r["HR_hi"], "P": r["p_wald"], "P (LR)": r["p_LR"],
                 "C covariates (in-sample)": r["C_cov_insample"], "C covariates + risk (in-sample)": r["C_cov_plus_risk_insample"]})
    for k, v in r.items():
        if k.startswith("type::"): rows.append({**base, "Term": f"{k[6:]} contribution (per SD)", "HR adjusted": v["HR"], "CI low": v["lo"], "CI high": v["hi"], "P": v["p"]})
S["S5_Clinical_adjusted_Cox"] = pd.DataFrame(rows)
S["S6_TCGA_phenotype_assoc"] = pd.read_csv(MR / "exp2/PHENO_ASSOC.csv")
ts = json.load(open(MR / "exp2/TP53_STATE.json")); rows = []
for o, v in ts.items():
    for sc, r in v.items():
        for meas in ("comp_weighted", "state_only"):
            rows.append({"Cancer": o.capitalize(), "Patients": sc, "Measure": {"comp_weighted": "Composition-weighted", "state_only": "State only (post hoc)"}[meas],
                         "TP53 mutant": r["n_mut"], "TP53 wild type": r["n_wt"], "AUC (mutant > wild type)": r[meas]["auc"], "Median difference (SD units)": r[meas]["median_diff"], "P (one-sided)": r[meas]["p_one_sided"]})
        ols = r["ols_state_only"]; rows.append({"Cancer": o.capitalize(), "Patients": sc, "Measure": "OLS state only ~ TP53 + malignant fraction (+PAM50 breast)", "TP53 mutant": r["n_mut"], "TP53 wild type": r["n_wt"],
                                                  "Beta TP53": ols["beta_TP53"], "CI low": ols["ci"][0], "CI high": ols["ci"][1], "P (two-sided)": ols["p_two_sided"]})
S["S7_TP53_proliferation"] = pd.DataFrame(rows)
dl = []
for o, on in ORG:
    sm = json.load(open(MR / f"exp1/{o}_SUMMARY.json"))
    for _, r in pd.read_csv(MR / f"exp1/{o}_PROGRAMS.csv").iterrows(): dl.append({"Cancer": on, "Unit": "program axis", "Name": r.program, "Contribution magnitude": r.magnitude, "ΔC (mean replacement)": r.dC_mean, "ΔC (permutation)": r.dC_perm})
    for _, r in pd.read_csv(MR / f"exp1/{o}_CELLTYPES.csv").iterrows(): dl.append({"Cancer": on, "Unit": "cell type", "Name": r.cell_type, "Contribution magnitude": r.magnitude, "ΔC (mean replacement)": r.dC_mean, "ΔC (permutation)": r.dC_perm})
    dl.append({"Cancer": on, "Unit": "top-5 axes jointly", "Name": ", ".join(sm["F2"]["top5"]), "ΔC (mean replacement)": sm["F2"]["dC_top5"], "CI low": sm["F2"]["dC_top5_ci95"][0], "CI high": sm["F2"]["dC_top5_ci95"][1],
               "Random 5-axis 95th percentile": sm["F2"]["random5_q95"], "P": sm["F2"]["p"], "Spearman magnitude vs ΔC (all axes)": sm["F1"]["spearman_magnitude_vs_dC_mean"]})
S["S8_Program_deletion"] = pd.DataFrame(dl)

S["S9_Risk_decomposition"] = pd.read_csv(ROOT / "results/mech/decomp/RISK_DECOMPOSITION.csv")
# ---------- S10 세포 유형·구획·Other 사전·CIBERSORT 대응 (레지스트리·train.type_mapping 그대로) ----------
import sys as _sys; _sys.path.insert(0, str(ROOT)); import train as _T
_LM22 = {"B": "B cells naive + memory", "PL": "Plasma cells", "T": "T cells CD8, CD4 naive, CD4 memory resting/activated, follicular helper, Tregs, gamma delta; NK cells resting/activated",
         "MY": "Monocytes, macrophages M0/M1/M2, dendritic cells resting/activated, mast cells resting/activated, eosinophils, neutrophils"}
_MAP = {"breast": {"B-cells": _LM22["B"], "Plasmablasts": _LM22["PL"], "T-cells": _LM22["T"], "Myeloid": _LM22["MY"]},
        "liver": {"B cells": _LM22["B"] + "; " + _LM22["PL"], "T cells": _LM22["T"], "TAMs": _LM22["MY"]},
        "pancreas": {"Immune": "all LM22 leukocyte classes"}}
_REF = {"breast": "GSE176078", "liver": "GSE151530", "pancreas": "GSE202051 (untreated)"}
_rows = []
for _o in ("breast", "liver", "pancreas"):
    _b = _T.load(_T.ORGANS[_o]["registry"]); _m = _T.type_mapping(_T.ORGANS[_o], _b); _n = list(_b["cell_type_names"])
    _D = _T.load(_T.ORGANS[_o]["stage1_data"], mmap=True); _ct = np.asarray(_D["sc_cell_type"]); _don = np.asarray(_D["sc_donor"]); _tr = set(int(x) for x in _D["donor_split"]["train"])
    _pri = dict(zip(_m["other_indices"], _m["other_prior"].tolist()))
    for _i, _t in enumerate(_n):
        _comp = "Tumor" if _i == _m["tumor_index"] else ("Immune" if _i in _m["immune_indices"] else "Other")
        _rows.append({"Cancer": _o.capitalize(), "Reference": _REF[_o], "Cell type": _t, "Compartment": _comp,
                      "Reference cells": int((_ct == _i).sum()), "Training-donor cells": int(((_ct == _i) & np.isin(_don, list(_tr))).sum()),
                      "Fixed prior within Other": round(_pri[_i], 4) if _i in _pri else None,
                      "CIBERSORT LM22 classes (immune head teacher)": _MAP[_o].get(_t, "")})
S["S10_Cell_types_compartments"] = pd.DataFrame(_rows)

TITLES = {"S1_Cindex_by_endpoint": "Table S1. C-index by cohort and endpoint for all models, including simplified bulk models (95% CI, 2,000 patient bootstrap), and paired differences, CellFactor minus each model, with 95% CI and two-sided bootstrap P (same resamples). CellFactor, ensemble of chains c1–c6. GSE10143 and GSE54236 are supplementary liver cohorts (not primary endpoints); for GSE54236 (OS*), all patients are treated as events because event status is unavailable.",
          "S2_Cohorts_datasets": "Table S2. Cohorts and datasets: single-cell references, TCGA training/held-out, external bulk cohorts (platform, n, events, selected-panel coverage) and spatial datasets.",
          "S3_Program_dictionary": "Table S3. Transcriptional program dictionary: source, gene counts, panel retention (≥3 genes and within-type r ≥ 0.5), genes and, for curated programs, supporting references. Curated programs are short lists of canonical markers selected from the cited literature; they are not copies of published signatures. CS, curated common cell state.",
          "S4_Selected_gene_panel": "Table S4. Selected gene panel (3,000 genes per cancer): inclusion reason, marker cell types and program membership.",
          "S5_Clinical_adjusted_Cox": "Table S5. Clinical-adjusted Cox models: CellFactor risk and top cell-type contributions.",
          "S6_TCGA_phenotype_assoc": "Table S6. Associations with independent TCGA phenotypes (TIL fraction, immune subtype, mutations); prespecified one-sided tests with Holm correction.",
          "S7_TP53_proliferation": "Table S7. TP53 mutation and malignant proliferation contribution (composition-weighted, prespecified; state only, post hoc).",
          "S8_Program_deletion": "Table S8. Program and cell-type deletion: contribution magnitude and decrease in external C-index (no retraining).",
          "S9_Risk_decomposition": "Table S9. Decomposition of CellFactor risk: full model, composition fixed at the training mean (state-only), composition-deviation term and composition only (95% CI, 2,000 patient bootstrap).",
          "S10_Cell_types_compartments": "Table S10. Reference cell types, their assignment to tumor, immune and other compartments, the fixed prior that splits the other compartment (mean composition of training-donor pseudo-bulk bags) and the CIBERSORT LM22 classes used to supervise the immune head, together with the number of high-risk cells (top 10% of reference cells by CellFactor risk, ranked within each cancer type), each cell type\u2019s share of the high-risk cells, and the observed/expected ratio (share of high-risk cells divided by share of all reference cells; 1 = proportional to abundance)."}
with pd.ExcelWriter(OUT / "Supplementary_Tables.xlsx") as W:
    pd.DataFrame({"Sheet": list(TITLES), "Title": list(TITLES.values())}).to_excel(W, sheet_name="Contents", index=False)
    for k, df in S.items():
        df.to_excel(W, sheet_name=k[:31], index=False); df.to_csv(OUT / f"tables/{k}.csv", index=False)


# ---- S10 I–K: 고위험 세포 수 · 고위험 내 비율 · O/E (원고 측 편집 2026-10-11 을 그대로 옮김; 입력 = final_results/01_umap/{org}/cell_type_counts.csv)
import csv, copy
from openpyxl import load_workbook


def add_s10_oe(wb, umap_dir):
    ws = wb['S10_Cell_types_compartments']
    cnt = {}
    for org, cancer in [('breast', 'Breast'), ('liver', 'Liver'), ('pancreas', 'Pancreas')]:
        R = list(csv.DictReader(open(f'{umap_dir}/{org}/cell_type_counts.csv')))
        N = sum(int(r['total_cells']) for r in R); H = sum(int(r['high_risk_cells']) for r in R)
        for r in R: cnt[(cancer, r['cell_type'])] = (int(r['total_cells']), int(r['high_risk_cells']), N, H)
    h0 = ws['H1']
    for col, name in zip('IJK', ['High-risk cells (top 10% by CellFactor risk)', 'Share of high-risk cells', 'Observed/expected ratio']):
        c = ws[f'{col}1']; c.value = name
        c.font, c.alignment, c.border, c.fill = (copy.copy(h0.font), copy.copy(h0.alignment), copy.copy(h0.border), copy.copy(h0.fill))
    for r in range(2, ws.max_row + 1):
        n, h, N, H = cnt[(ws[f'A{r}'].value, ws[f'C{r}'].value)]
        assert ws[f'E{r}'].value == n          # Reference cells must match the UMAP counts
        e = ws[f'E{r}']
        for col, v, fmt in [('I', h, '#,##0'), ('J', h / H, '0.0%'), ('K', (h / H) / (n / N), '0.00')]:
            c = ws[f'{col}{r}']; c.value = v; c.number_format = fmt
            c.font, c.alignment, c.border = copy.copy(e.font), copy.copy(e.alignment), copy.copy(e.border)
    for col, w in zip('IJK', [22, 14, 14]): ws.column_dimensions[col].width = w


import openpyxl, sys as _sys
REFS = Path(__file__).resolve().parent / "S3_PROGRAM_REFERENCES.json"   # {"R": {키: 서지}, "P": {프로그램: [키, ...]}} — 원고 측 xl/s3_refs.py 의 R·P 사전 그대로


def add_s3_refs(wb, R, P):
    ws = wb['S3_Program_dictionary']; hdr = [c.value for c in ws[1]]; assert hdr[-1] == 'Genes'
    col = len(hdr) + 1; h0 = ws.cell(1, len(hdr))
    c = ws.cell(1, col, 'Supporting references (curated programs)')
    c.font, c.alignment, c.fill, c.border = copy.copy(h0.font), copy.copy(h0.alignment), copy.copy(h0.fill), copy.copy(h0.border)
    seen = set()
    for r in range(2, ws.max_row + 1):
        prog, src = ws.cell(r, 2).value, str(ws.cell(r, 3).value or '')
        if not src.startswith('Curated'): continue
        ref = ws.cell(r, len(hdr)); c = ws.cell(r, col, '; '.join(R[k] for k in P[prog]))
        c.font, c.alignment = copy.copy(ref.font), copy.copy(ref.alignment); seen.add(prog)
    assert seen == set(P), set(P) - seen
    ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 80


_wb = load_workbook(OUT / "Supplementary_Tables.xlsx"); add_s10_oe(_wb, str(ROOT / "final_results/01_umap"))
if REFS.exists():
    _rp = json.load(open(REFS)); add_s3_refs(_wb, _rp["R"], _rp["P"])
else:
    print(f"[warn] {REFS} 없음 — S3 J열(근거 문헌) 미포함", file=_sys.stderr)
_wb.save(OUT / "Supplementary_Tables.xlsx")
_ws3 = _wb["S3_Program_dictionary"]
pd.DataFrame([[c.value for c in row] for row in _ws3.iter_rows(min_row=2)], columns=[c.value for c in _ws3[1]]).to_csv(OUT / "tables/S3_Program_dictionary.csv", index=False)
_ws = _wb["S10_Cell_types_compartments"]   # csv 사본도 같은 값으로
pd.DataFrame([[c.value for c in row] for row in _ws.iter_rows(min_row=2)], columns=[c.value for c in _ws[1]]).to_csv(OUT / "tables/S10_Cell_types_compartments.csv", index=False)
json.dump(TITLES, open(OUT / "tables/TITLES.json", "w"), indent=1)
print({k: len(v) for k, v in S.items()}); print("coverage", {k: round(v, 3) for k, v in cov.items()})
