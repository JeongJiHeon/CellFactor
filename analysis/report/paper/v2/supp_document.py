#!/usr/bin/env python3
"""Supplementary Information 문서(.docx) — 그림 S1–S7 + 범례, 보충 표 S1–S10 목록(본체는 Supplementary_Tables.xlsx).
출력 final_results/figures/manuscript_v2/Supp/Supplementary_Information.docx"""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
ROOT = Path(__file__).resolve().parents[4]; S = ROOT / "final_results/figures/manuscript_v2/Supp"
FIGS = [
    ("S1_KM_all.png", "Figure S1. Kaplan–Meier survival stratification for CellFactor and comparator methods.",
     "Overall survival in the largest external cohort of each cancer (rows: breast SCAN-B and METABRIC, liver GSE14520, pancreas GSE62452) for five models (columns). "
     "Patients were split into the top 20% versus the remaining 80% of each model's own risk score. Tick marks, censoring. P, two-sided log-rank test (bold, P < 0.05). "
     "Main Figure 2f summarizes the CellFactor P values."),
    ("S2_representation.png", "Figure S2. Characterization of the CellFactor state representation.",
     "(a) Named versus free axes. Within-cell-type cluster read-out using only named program axes or only the four free axes, expressed as retention = (AUC − 0.5)/(AUC_all − 0.5). "
     "Numbers, seeds (of 10) in which four randomly chosen named axes outperformed the four free axes (descriptive counts, not a statistical test). Liver T cells are shown for the prespecified analysis including all Leiden clusters "
     "and for a post hoc sensitivity analysis excluding clusters with canonical hepatocyte genes among their top three markers (17 of 41 cluster instances across ten seeds; two recurrent clusters). "
     "CellFactor mean AUROC was 0.711 with all liver T-cell clusters and 0.828 after post hoc exclusion of hepatocyte-marker-associated clusters. "
     "Settings: internal held-out donors of GSE151530 (liver), GSE202051 (pancreas) and GSE176078 (breast), and the independent breast cohort GSE161529 (external). "
     "(b) Read-out direction. For each Leiden cluster, Spearman ρ between the read-out contribution of each axis (W × SD of z′) and the model-independent change in the corresponding program score (Δ) "
     "in that cluster (blue), compared with the mean over other clusters (grey, null). ρ̃, median; cl., number of clusters. "
     "(c–e) Best-case clusters (highest ρ per cancer), shown as examples rather than representative cases; the median ρ over all clusters is given in b. "
     "The liver example is a proliferating T-cell cluster and is not one of the hepatocyte-marker-associated clusters. Circled points, three axes with the largest read-out contribution."),
    ("S3_objective_ablation.png", "Figure S3. Training-objective ablation of the bulk state estimator.",
     "The single-cell representation z′ was fixed and only the bulk state estimator H was retrained without the paired-delta loss, without the CORAL alignment loss, or without both. "
     "(a) Difference in within-cell-type cluster read-out AUROC, CellFactor minus ablated model (× 100; seed-paired mean and 95% CI over 10 seeds). "
     "(b) Median read-out direction ρ for each training objective. All liver Leiden clusters are included (prespecified analysis)."),
    ("S4_literature_direction.png", "Figure S4. Concordance of program weight directions with literature expectations.",
     "Fraction of pre-registered named axes whose prognostic weight sign matched the direction expected from the literature (blue), "
     "against a permutation null (200 permutations; grey, mean ± 1.96 SD). Numbers, concordant/tested axes and permutation P. "
     "No cancer reached P < 0.05; weight directions describe the fitted model and are not evidence of biological effect."),
    ("S5_program_deletion.png", "Figure S5. Predictive reliance on program axes.",
     "Decrease in the mean external C-index after replacing the five program axes with the largest contribution magnitude by their cohort means "
     "(blue line; bar, patient-bootstrap 95% CI), compared with 1,000 random sets of five axes (grey; dashed line, 95th percentile). "
     "Axes were ranked by contribution magnitude, which does not use outcomes. No retraining was performed, so the analysis measures the reliance of the trained model on these axes, "
     "not whether other axes could recover the information. A clear effect was observed only in breast cancer; across individual axes, contribution magnitude did not track the deletion effect "
     "in any cancer (Spearman ρ −0.05, −0.02 and 0.20; Table S8)."),
    ("S6_domain_gap.png", "Figure S6. Pseudo-bulk–bulk domain gap and the effect of conditional CORAL.",
     "(a) PCA of the 3,000-gene selected panel for training-donor pseudo-bulk bags (grey, n = 300) and TCGA bulk (blue), in a common unit "
     "(log2 CPM re-normalised within the same gene set) and as the domain-standardised model input. AUROC, 5-fold cross-validated domain classifiers "
     "(LR, L2 logistic regression; RF, random forest); for the model input, LR uses domain standardisation re-estimated within training folds, because standardising "
     "on the full data sets each domain mean to exactly zero and biases cross-validated AUROC below 0.5. Per-domain standardisation removes linear separability, but RF still separates the domains. "
     "The domains also differ in patients and cell-type composition, so separation is not attributable to technical batch effects alone. "
     "(b) Gene-wise mean expression (common unit) and per-gene detection rate; zero fractions (pseudo-bulk vs TCGA) are shown above. Means are compared across different individuals and do not measure per-patient agreement. "
     "(c) State-estimator bottleneck features after regressing out estimated composition (pooled over both domains): PCA fitted on pseudo-bulk only, with held-out TCGA samples projected; "
     "axes in pseudo-bulk standard deviations and identical limits for the two models (chain c1). The two models share Stage 1′, bags and training seeds and differ only in the CORAL weight. "
     "(d) Discrepancy for each chain (c1–c6; lines join the same chain). The CORAL objective residualises each domain separately, so its mean term is zero by construction and it penalises only covariance differences. "
     "With CORAL, the composition-adjusted mean difference halved in feature units, but the feature scale shrank more, so the mean difference in standard-deviation units increased. "
     "(e) External C-index of the c1–c6 ensemble with patient-bootstrap 95% CIs (Δ, paired difference CORAL on − off; 2,000 resamples); "
     "relative change in held-out-donor pseudo-bulk state error (MAE relative to a constant predictor; open circles, breast, whose donors were all used in training, so the breast value is in-sample); "
     "and domain-classifier AUROC (median of chains) on composition-adjusted features and after removing domain means within training folds. "
     "Linear separation disappears once means are removed, whereas a random forest still separates the domains; a control splitting pseudo-bulk bags at random gave AUROC ≈ 0.5. "
     "CORAL therefore reduced the covariance mismatch it targets without making the domains indistinguishable; state error was similar or lower, and the external C-index was 0.010–0.016 higher, "
     "with a CI excluding zero only in breast cancer."),
    ("S7_simulation.png", "Figure S7. Recovery of known risk contributions in semi-synthetic cohorts.",
     "Synthetic breast cancer patients (1,024 cells each; 500 training and 500 test patients per replicate; 10 replicates) were assembled from GSE161529 single cells, "
     "which were not used to train CellFactor; training and test patients were drawn from disjoint halves of the 52 donors. Composition (Dirichlet) and cell states "
     "(cells sampled with weights tilted towards patient-specific program scores) were varied independently. True risk was a pre-specified sum over cell types of "
     "composition × program state (malignant proliferation +, malignant early estrogen response −, T-cell cytotoxicity −, CAF TGF-β +; endothelial hypoxia modulated with zero effect), "
     "with true states defined as gene-set scores of the sampled cells, independent of model coordinates. Survival was exponential with 40% events. "
     "Scenarios: A, independent; B, malignant fraction correlated with the proliferation latent (ρ = 0.6); C, malignant hypoxia (true effect) and endothelial hypoxia (no effect) sharing a latent factor (ρ = 0, 0.5, 0.9); "
     "D, scenario A plus a gene-wise shift (TCGA minus synthetic mean) and noise. CellFactor used the paper's frozen composition and state estimators (chains c1–c6); only the risk readout was refitted (α free, as in the reported model). "
     "Oracles replace estimated composition, estimated states (M*, mean cell-level z′ of the sampled cells), or both with the truth. "
     "(a) Spearman correlation between estimated and true cell-type contributions in test patients (dots, replicates; bar, median; dotted line, 0.5). "
     "(b) Spearman correlation for each true program contribution. (c) Number of replicates in which the sign of the CellFactor program effect matched the truth. "
     "(d) Share of the hypoxia contribution assigned to endothelial cells (true share 0). (e) Test C-index of the true risk, the oracle with both inputs, CellFactor and bulk ridge Cox. "
     "(f) Spearman correlation between true states and bulk-estimated states (M̂) or cell-level states (M*), scenario A. "
     "CellFactor assigned most risk variance to the cell types carrying effects (median share 0.90–0.94) and recovered the malignant contribution (median ρ 0.74–0.80) and the proliferation direction (9–10/10), "
     "but not the patient-level T-cell or CAF contributions (median ρ −0.03 to 0.12) or the malignant hypoxia effect; with true cell states these reached 0.74–0.83. "
     "The limiting step is therefore state estimation from bulk for minor cell types, not the representation or the risk readout. "
     "CAF and endothelial fractions are not estimated individually (fixed split of the 'Other' compartment)."),
]
TIT = json.load(open(S / "tables/TITLES.json"))
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10)
doc.add_heading("Supplementary Information", 0)
doc.add_paragraph("CellFactor: cell-type-resolved prognostic states from bulk transcriptomes")
doc.add_heading("Contents", 1)
for _, t, _ in FIGS: doc.add_paragraph(t, style="List Bullet")
for k, t in TIT.items(): doc.add_paragraph(t.split(":")[0].rstrip(".") + ".", style="List Bullet")
doc.add_heading("Supplementary Figures", 1)
for f, t, leg in FIGS:
    doc.add_page_break(); doc.add_picture(str(S / f), width=Inches(6.5))
    p = doc.add_paragraph(); r = p.add_run(t + " "); r.bold = True; p.add_run(leg)
doc.add_page_break(); doc.add_heading("Supplementary Tables", 1)
doc.add_paragraph("Supplementary Tables S1–S10 are provided as Supplementary_Tables.xlsx (one sheet per table).")
for k, t in TIT.items():
    p = doc.add_paragraph(); r = p.add_run(t.split(". ", 1)[0] + ". "); r.bold = True; p.add_run(t.split(". ", 1)[1] + f" (sheet {k})")
doc.save(S / "Supplementary_Information.docx"); print("ok")
